from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping

from futonhub.cloud.audit import AuditEvent, OperationSnapshot, new_operation_id, write_audit_event, write_snapshot
from futonhub.core.catalog_policy import is_discontinued_commercial_status
from gestorwoo.config import Settings, load_settings


PROCESS_ROTATION_C = "rotation_c"
PROCESS_STOCK = "stock"
PROCESS_SUPPLIER_PRICES = "supplier_prices"

READY = "READY"
NO_CHANGE = "NO_CHANGE"
INVALID_ID = "INVALID_ID"
INVALID_VALUE = "INVALID_VALUE"
DUPLICATE_ID = "DUPLICATE_ID"
DESCATALOGADO_NO_ACTUALIZABLE = "DESCATALOGADO_NO_ACTUALIZABLE"
PRIMARY_SUPPLIER_UNRESOLVED = "PRIMARY_SUPPLIER_UNRESOLVED"
SUPPLIER_PRICES_READ_ERROR = "SUPPLIER_PRICES_READ_ERROR"

BLOCKING_STATUSES = {
    INVALID_ID,
    INVALID_VALUE,
    DUPLICATE_ID,
    DESCATALOGADO_NO_ACTUALIZABLE,
    PRIMARY_SUPPLIER_UNRESOLVED,
    SUPPLIER_PRICES_READ_ERROR,
}

PROCESS_DEFINITIONS: dict[str, dict[str, Any]] = {
    PROCESS_ROTATION_C: {
        "label": "Rotación C",
        "template_name": "Plantilla_RotacionC.xlsx",
        "headers": ("ID", "Rotación C"),
    },
    PROCESS_STOCK: {
        "label": "Stock",
        "template_name": "Plantilla_Stock.xlsx",
        "headers": ("ID", "Stock Tienda", "Stock Warehouse"),
    },
    PROCESS_SUPPLIER_PRICES: {
        "label": "Precios Proveedores",
        "template_name": "Plantilla_Precios_Proveedores.xlsx",
        "headers": ("ID", "Precio Principal", "Pascal"),
    },
}

INVENTORY_UPDATE_SELECT_COLUMNS = (
    "item_id,name,hub_item_code,heca_reference,woo_sku,item_record_type,base_item_code,is_pack,"
    "commercial_status,rotation_c,store_stock,warehouse_stock,weighted_average_cost,"
    "primary_supplier_price,pascal_price,supplier_order_provider,source_row,updated_at"
)

SUPPLIER_PRICE_SELECT_COLUMNS = "item_id,supplier,price,currency,source,updated_at"
PRIMARY_SUPPLIERS = {"Ekomat", "Cipta", "Hemei"}


class UpdatesValidationError(ValueError):
    pass


class SupplierPricesReadError(UpdatesValidationError):
    pass


class UpdateApplyError(RuntimeError):
    def __init__(self, message: str, result: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.result = result or {}


@dataclass(frozen=True)
class ParsedUpdateRow:
    row_number: int
    values: dict[str, Any]


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _json_safe(value: Any) -> Any:
    import json

    try:
        return json.loads(json.dumps(value, ensure_ascii=False, default=str))
    except Exception:
        return {"_raw": str(value)}


def _decimal_or_none(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return Decimal(text.replace(",", "."))
    except InvalidOperation:
        return None


def _parse_decimal(value: Any, *, allow_blank: bool, field_label: str, allow_negative: bool = False) -> tuple[Decimal | None, str]:
    if _is_blank(value):
        return (None, "") if allow_blank else (None, f"{field_label} vacio")
    number = _decimal_or_none(value)
    if number is None:
        return None, f"{field_label} no es numerico"
    if not allow_negative and number < 0:
        return None, f"{field_label} no puede ser negativo"
    return number, ""


def _decimal_equal(left: Any, right: Any) -> bool:
    left_number = _decimal_or_none(left)
    right_number = _decimal_or_none(right)
    if left_number is None or right_number is None:
        return left_number is None and right_number is None
    return left_number == right_number


def _decimal_to_payload(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value.quantize(Decimal("0.0001")))


def _source_row_dict(row: Mapping[str, Any] | None) -> dict[str, Any]:
    if not isinstance(row, Mapping):
        return {}
    value = row.get("source_row")
    if isinstance(value, dict):
        return value
    if isinstance(value, str) and value.strip():
        import json

        try:
            decoded = json.loads(value)
        except Exception:
            return {}
        return decoded if isinstance(decoded, dict) else {}
    return {}


def _normalize_supplier(value: Any) -> str:
    text = str(value or "").strip().casefold()
    aliases = {
        "ekomat": "Ekomat",
        "cipta": "Cipta",
        "hemei": "Hemei",
        "heimei": "Hemei",
        "pascal": "Pascal",
    }
    return aliases.get(text, "")


def _canonical_price_value(value: Any) -> float | None:
    number = _decimal_or_none(value)
    return _decimal_to_payload(number) if number is not None else None


def _row_matches_requested_id(row: Mapping[str, Any], requested_id: str) -> bool:
    requested = str(requested_id or "").strip()
    for field in ("heca_reference", "hub_item_code", "woo_sku"):
        if str(row.get(field) or "").strip() == requested:
            return True
    return False


def _is_simple_inventory_row(row: Mapping[str, Any]) -> bool:
    record_type = str(row.get("item_record_type") or "").strip().lower()
    if record_type not in {"", "simple"}:
        return False
    if str(row.get("base_item_code") or "").strip():
        return False
    return True


def _select_rows_by_exact_value(session, column: str, value: Any) -> list[dict[str, Any]]:
    response = (
        session.client.table("inventory_items")
        .select(INVENTORY_UPDATE_SELECT_COLUMNS)
        .eq(column, value)
        .limit(3)
        .execute()
    )
    return [dict(row) for row in (getattr(response, "data", None) or [])]


def resolve_update_inventory_item(session, requested_id: Any) -> tuple[dict[str, Any] | None, str]:
    code = str(requested_id or "").strip()
    if not code:
        return None, "ID vacio"

    found: dict[int, dict[str, Any]] = {}
    diagnostics: list[str] = []
    for column in ("heca_reference", "hub_item_code", "woo_sku"):
        try:
            rows = _select_rows_by_exact_value(session, column, code)
        except Exception as exc:
            diagnostics.append(f"{column}: {exc}")
            continue
        for row in rows:
            try:
                found[int(row.get("item_id"))] = row
            except Exception:
                diagnostics.append(f"{column}: item_id invalido")

    if len(found) == 1:
        row = next(iter(found.values()))
        if not _row_matches_requested_id(row, code):
            return None, "La fila encontrada no conserva el ID literal solicitado"
        if not _is_simple_inventory_row(row):
            return None, "La fila no es un registro fisico simple actualizable"
        return row, ""
    if len(found) > 1:
        return None, "Resolucion exacta ambigua para el ID"
    if diagnostics:
        return None, "; ".join(diagnostics)
    return None, "No existe coincidencia exacta"


def _fetch_inventory_item_by_id(session, item_id: int | str) -> dict[str, Any] | None:
    response = (
        session.client.table("inventory_items")
        .select(INVENTORY_UPDATE_SELECT_COLUMNS)
        .eq("item_id", int(item_id))
        .limit(1)
        .execute()
    )
    rows = getattr(response, "data", None) or []
    return dict(rows[0]) if rows else None


def _fetch_supplier_prices_for_item(session, item_id: int | str) -> list[dict[str, Any]]:
    try:
        response = (
            session.client.table("supplier_prices")
            .select(SUPPLIER_PRICE_SELECT_COLUMNS)
            .eq("item_id", int(item_id))
            .execute()
        )
    except Exception as exc:
        raise SupplierPricesReadError(f"No se pudo leer supplier_prices para item_id {item_id}: {exc}") from exc
    return [dict(row) for row in (getattr(response, "data", None) or [])]


def _supplier_price_map(rows: Iterable[Mapping[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        supplier = _normalize_supplier(row.get("supplier"))
        if supplier:
            result[supplier] = dict(row)
    return result


def resolve_primary_supplier(row: Mapping[str, Any], supplier_price_rows: Iterable[Mapping[str, Any]] | None = None) -> str | None:
    source = _source_row_dict(row)
    candidates: list[str] = []

    for key in ("primary_supplier", "supplier_price_provider", "supplier_order_provider"):
        supplier = _normalize_supplier(source.get(key))
        if supplier in PRIMARY_SUPPLIERS:
            candidates.append(supplier)

    migration = source.get("supplier_price_migration")
    if isinstance(migration, dict):
        for entry in migration.get("sources") or []:
            if isinstance(entry, Mapping):
                supplier = _normalize_supplier(entry.get("supplier"))
                if supplier in PRIMARY_SUPPLIERS:
                    candidates.append(supplier)

    row_supplier = _normalize_supplier(row.get("supplier_order_provider"))
    if row_supplier in PRIMARY_SUPPLIERS:
        candidates.append(row_supplier)

    for supplier in candidates:
        return supplier

    suppliers_from_table = {
        _normalize_supplier(price_row.get("supplier"))
        for price_row in (supplier_price_rows or [])
        if _normalize_supplier(price_row.get("supplier")) in PRIMARY_SUPPLIERS
    }
    if len(suppliers_from_table) == 1:
        return next(iter(suppliers_from_table))
    return None


def _excel_id_value(cell: Any) -> str:
    value = getattr(cell, "value", cell)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    else:
        raw = str(value).strip()
    number_format = str(getattr(cell, "number_format", "") or "")
    if raw.isdigit() and set(number_format) <= {"0"} and len(number_format) > len(raw):
        return raw.zfill(len(number_format))
    return raw


def _trim_header_values(values: list[Any]) -> list[str]:
    result = ["" if value is None else str(value) for value in values]
    while result and result[-1] == "":
        result.pop()
    return result


def _definition(process: str) -> dict[str, Any]:
    if process not in PROCESS_DEFINITIONS:
        raise UpdatesValidationError(f"Proceso de actualizacion no soportado: {process}")
    return PROCESS_DEFINITIONS[process]


def read_updates_workbook(path: str | Path, process: str) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    definition = _definition(process)
    expected_headers = list(definition["headers"])
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1))
        headers = _trim_header_values([cell.value for cell in header_cells])
        if headers != expected_headers:
            raise UpdatesValidationError(
                f"Headers invalidos para {definition['label']}. Esperado: {' | '.join(expected_headers)}"
            )

        rows: list[dict[str, Any]] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(expected_headers)), start=2):
            raw_values = [cell.value for cell in cells]
            if all(_is_blank(value) for value in raw_values):
                continue
            values: dict[str, Any] = {"_row_number": row_number}
            for header, cell in zip(expected_headers, cells):
                values[header] = _excel_id_value(cell) if header == "ID" else cell.value
            rows.append(values)
        return rows
    finally:
        workbook.close()


def write_update_template(process: str, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    definition = _definition(process)
    output_path = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = definition["label"][:31]
    for index, header in enumerate(definition["headers"], start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = max(18, len(header) + 4)
    sheet.freeze_panes = "A2"
    workbook.save(output_path)
    return output_path


def _error_row(raw: Mapping[str, Any], status: str, reason: str, process: str, requested_id: str | None = None) -> dict[str, Any]:
    return {
        "process": process,
        "row_number": raw.get("_row_number"),
        "id": requested_id if requested_id is not None else str(raw.get("ID") or "").strip(),
        "item_id": None,
        "name": "",
        "status": status,
        "reason": reason,
        "changes": [],
    }


def _base_preview_row(process: str, raw: Mapping[str, Any], item: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "process": process,
        "row_number": raw.get("_row_number"),
        "id": str(raw.get("ID") or "").strip(),
        "item_id": int(item.get("item_id")),
        "name": item.get("name") or "",
        "heca_reference": item.get("heca_reference") or "",
        "hub_item_code": item.get("hub_item_code") or "",
        "commercial_status": item.get("commercial_status") or "",
        "status": READY,
        "reason": "",
        "changes": [],
    }


def _change(
    *,
    field: str,
    current: Any,
    new_value: Decimal,
    supplier: str | None = None,
    supplier_current: Any = None,
) -> dict[str, Any]:
    new_payload = _decimal_to_payload(new_value)
    return {
        "field": field,
        "current": current,
        "new": new_payload,
        "expected_current": current,
        "supplier": supplier,
        "supplier_current": supplier_current,
        "expected_supplier_current": supplier_current,
        "inventory_changed": not _decimal_equal(current, new_payload),
        "supplier_price_changed": supplier is not None and not _decimal_equal(supplier_current, new_payload),
    }


def _preview_rotation(session, raw: Mapping[str, Any], item: Mapping[str, Any]) -> dict[str, Any]:
    row = _base_preview_row(PROCESS_ROTATION_C, raw, item)
    value, error = _parse_decimal(raw.get("Rotación C"), allow_blank=False, field_label="Rotacion C", allow_negative=True)
    row["rotation_current"] = item.get("rotation_c")
    row["rotation_new"] = _decimal_to_payload(value) if value is not None else None
    if error:
        row.update({"status": INVALID_VALUE, "reason": error})
        return row
    change = _change(field="rotation_c", current=item.get("rotation_c"), new_value=value)
    if not change["inventory_changed"]:
        row.update({"status": NO_CHANGE, "reason": "Sin cambios"})
    else:
        row["changes"] = [change]
    return row


def _preview_stock(session, raw: Mapping[str, Any], item: Mapping[str, Any]) -> dict[str, Any]:
    row = _base_preview_row(PROCESS_STOCK, raw, item)
    store, store_error = _parse_decimal(raw.get("Stock Tienda"), allow_blank=False, field_label="Stock Tienda")
    warehouse, warehouse_error = _parse_decimal(raw.get("Stock Warehouse"), allow_blank=False, field_label="Stock Warehouse")
    row.update(
        {
            "store_stock_current": item.get("store_stock"),
            "store_stock_new": _decimal_to_payload(store) if store is not None else None,
            "warehouse_stock_current": item.get("warehouse_stock"),
            "warehouse_stock_new": _decimal_to_payload(warehouse) if warehouse is not None else None,
            "stock_total_current": (_canonical_price_value(item.get("store_stock")) or 0.0)
            + (_canonical_price_value(item.get("warehouse_stock")) or 0.0),
        }
    )
    if store_error or warehouse_error:
        row.update({"status": INVALID_VALUE, "reason": "; ".join(error for error in (store_error, warehouse_error) if error)})
        return row
    row["stock_total_new"] = float(store or Decimal("0")) + float(warehouse or Decimal("0"))
    changes = [
        _change(field="store_stock", current=item.get("store_stock"), new_value=store),
        _change(field="warehouse_stock", current=item.get("warehouse_stock"), new_value=warehouse),
    ]
    changes = [change for change in changes if change["inventory_changed"]]
    if not changes:
        row.update({"status": NO_CHANGE, "reason": "Sin cambios"})
    else:
        row["changes"] = changes
    return row


def _preview_supplier_prices(session, raw: Mapping[str, Any], item: Mapping[str, Any]) -> dict[str, Any]:
    row = _base_preview_row(PROCESS_SUPPLIER_PRICES, raw, item)
    try:
        supplier_rows = _fetch_supplier_prices_for_item(session, item.get("item_id"))
    except SupplierPricesReadError as exc:
        row.update({"status": SUPPLIER_PRICES_READ_ERROR, "reason": str(exc)})
        return row
    supplier_map = _supplier_price_map(supplier_rows)
    primary_supplier = resolve_primary_supplier(item, supplier_rows)
    row.update(
        {
            "primary_supplier": primary_supplier or "",
            "primary_supplier_price_current": item.get("primary_supplier_price"),
            "primary_supplier_price_new": "",
            "pascal_price_current": item.get("pascal_price"),
            "pascal_price_new": "",
        }
    )

    changes: list[dict[str, Any]] = []
    errors: list[str] = []
    primary_raw = raw.get("Precio Principal")
    if not _is_blank(primary_raw):
        value, error = _parse_decimal(primary_raw, allow_blank=True, field_label="Precio Principal")
        row["primary_supplier_price_new"] = _decimal_to_payload(value) if value is not None else ""
        if error:
            errors.append(error)
        elif not primary_supplier:
            row.update({"status": PRIMARY_SUPPLIER_UNRESOLVED, "reason": "Proveedor principal no resuelto"})
            return row
        else:
            current_supplier_row = supplier_map.get(primary_supplier) or {}
            changes.append(
                _change(
                    field="primary_supplier_price",
                    current=item.get("primary_supplier_price"),
                    new_value=value,
                    supplier=primary_supplier,
                    supplier_current=current_supplier_row.get("price"),
                )
            )

    pascal_raw = raw.get("Pascal")
    if not _is_blank(pascal_raw):
        value, error = _parse_decimal(pascal_raw, allow_blank=True, field_label="Pascal")
        row["pascal_price_new"] = _decimal_to_payload(value) if value is not None else ""
        if error:
            errors.append(error)
        else:
            current_supplier_row = supplier_map.get("Pascal") or {}
            changes.append(
                _change(
                    field="pascal_price",
                    current=item.get("pascal_price"),
                    new_value=value,
                    supplier="Pascal",
                    supplier_current=current_supplier_row.get("price"),
                )
            )

    if errors:
        row.update({"status": INVALID_VALUE, "reason": "; ".join(errors)})
        return row
    changes = [change for change in changes if change["inventory_changed"] or change["supplier_price_changed"]]
    if not changes:
        row.update({"status": NO_CHANGE, "reason": "Sin cambios"})
    else:
        row["changes"] = changes
    return row


def preview_updates(session, process: str, rows: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    _definition(process)
    raw_rows = [dict(row) for row in rows]
    requested_ids = [str(row.get("ID") or "").strip() for row in raw_rows if str(row.get("ID") or "").strip()]
    duplicate_ids = {value for value in requested_ids if requested_ids.count(value) > 1}
    preview_rows: list[dict[str, Any]] = []

    for raw in raw_rows:
        requested_id = str(raw.get("ID") or "").strip()
        if not requested_id:
            preview_rows.append(_error_row(raw, INVALID_ID, "ID vacio", process, requested_id))
            continue
        if requested_id in duplicate_ids:
            preview_rows.append(_error_row(raw, DUPLICATE_ID, "ID duplicado en el Excel", process, requested_id))
            continue
        item, reason = resolve_update_inventory_item(session, requested_id)
        if item is None:
            preview_rows.append(_error_row(raw, INVALID_ID, reason or "No existe coincidencia exacta", process, requested_id))
            continue
        if is_discontinued_commercial_status(item.get("commercial_status")):
            row = _base_preview_row(process, raw, item)
            row.update(
                {
                    "status": DESCATALOGADO_NO_ACTUALIZABLE,
                    "reason": "commercial_status=Descatalogado no es actualizable desde Actualizaciones",
                }
            )
            preview_rows.append(row)
            continue
        if process == PROCESS_ROTATION_C:
            preview_rows.append(_preview_rotation(session, raw, item))
        elif process == PROCESS_STOCK:
            preview_rows.append(_preview_stock(session, raw, item))
        elif process == PROCESS_SUPPLIER_PRICES:
            preview_rows.append(_preview_supplier_prices(session, raw, item))
        else:
            raise UpdatesValidationError(f"Proceso de actualizacion no soportado: {process}")

    status_counts: dict[str, int] = {}
    for row in preview_rows:
        status_counts[row["status"]] = status_counts.get(row["status"], 0) + 1
    ready_rows = [row for row in preview_rows if row["status"] == READY]
    return {
        "process": process,
        "label": PROCESS_DEFINITIONS[process]["label"],
        "valid": not any(row["status"] in BLOCKING_STATUSES for row in preview_rows),
        "apply_enabled": bool(ready_rows) and not any(row["status"] in BLOCKING_STATUSES for row in preview_rows),
        "rows": preview_rows,
        "row_count": len(preview_rows),
        "ready_count": len(ready_rows),
        "no_change_count": status_counts.get(NO_CHANGE, 0),
        "status_counts": status_counts,
        "write_cell_count": sum(len(row.get("changes") or []) for row in ready_rows),
    }


def preview_updates_from_excel(session, process: str, path: str | Path) -> dict[str, Any]:
    rows = read_updates_workbook(path, process)
    preview = preview_updates(session, process, rows)
    preview["source_path"] = str(path)
    return preview


def _validate_live_row(preview_row: Mapping[str, Any], live: Mapping[str, Any], supplier_rows: list[dict[str, Any]]) -> None:
    if is_discontinued_commercial_status(live.get("commercial_status")):
        raise UpdateApplyError(f"{preview_row.get('id')}: el item paso a Descatalogado antes del apply")
    if not _is_simple_inventory_row(live):
        raise UpdateApplyError(f"{preview_row.get('id')}: la identidad dejo de ser registro simple")
    if not _row_matches_requested_id(live, str(preview_row.get("id") or "")):
        raise UpdateApplyError(f"{preview_row.get('id')}: la identidad live ya no coincide con el ID literal")
    if preview_row.get("process") == PROCESS_SUPPLIER_PRICES:
        live_primary = resolve_primary_supplier(live, supplier_rows)
        expected_primary = str(preview_row.get("primary_supplier") or "").strip()
        if expected_primary and live_primary != expected_primary:
            raise UpdateApplyError(f"{preview_row.get('id')}: drift de proveedor principal {expected_primary} -> {live_primary or 'sin resolver'}")
    supplier_map = _supplier_price_map(supplier_rows)
    for change in preview_row.get("changes") or []:
        field = str(change.get("field") or "")
        if change.get("inventory_changed") and not _decimal_equal(live.get(field), change.get("expected_current")):
            raise UpdateApplyError(f"{preview_row.get('id')}: drift en {field}")
        supplier = str(change.get("supplier") or "").strip()
        if supplier and change.get("supplier_price_changed"):
            supplier_row = supplier_map.get(supplier) or {}
            if not _decimal_equal(supplier_row.get("price"), change.get("expected_supplier_current")):
                raise UpdateApplyError(f"{preview_row.get('id')}: drift en supplier_prices[{supplier}]")


def _update_inventory_fields(session, item_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    response = session.client.table("inventory_items").update(payload).eq("item_id", int(item_id)).execute()
    rows = getattr(response, "data", None) or []
    return dict(rows[0]) if rows else {**payload, "item_id": int(item_id)}


def _upsert_supplier_price(session, item_id: int, supplier: str, price: Any, operation_id: str, now: str) -> dict[str, Any]:
    payload = {
        "item_id": int(item_id),
        "supplier": supplier,
        "price": price,
        "currency": "EUR",
        "source": f"updates_module:{operation_id}",
        "updated_at": now,
    }
    response = session.client.table("supplier_prices").upsert(payload, on_conflict="item_id,supplier").execute()
    rows = getattr(response, "data", None) or []
    return dict(rows[0]) if rows else payload


def _delete_supplier_price(session, item_id: int, supplier: str) -> None:
    session.client.table("supplier_prices").delete().eq("item_id", int(item_id)).eq("supplier", supplier).execute()


def _rollback_applied_rows(session, applied_rows: list[dict[str, Any]], *, now: str) -> dict[str, Any]:
    rollback_errors: list[str] = []
    restored = 0
    for applied in reversed(applied_rows):
        item_id = int(applied["item_id"])
        try:
            inventory_payload = {
                field: applied["before_inventory"].get(field)
                for field in applied.get("inventory_fields") or []
            }
            if inventory_payload:
                inventory_payload["updated_at"] = now
                inventory_payload["updated_by"] = getattr(session, "user_id", None)
                _update_inventory_fields(session, item_id, inventory_payload)
            before_supplier_map = _supplier_price_map(applied.get("before_supplier_prices") or [])
            for supplier in applied.get("supplier_fields") or []:
                before_supplier = before_supplier_map.get(supplier)
                if before_supplier:
                    _upsert_supplier_price(
                        session,
                        item_id,
                        supplier,
                        before_supplier.get("price"),
                        str(applied.get("operation_id") or "rollback"),
                        now,
                    )
                else:
                    _delete_supplier_price(session, item_id, supplier)
            restored += 1
        except Exception as exc:
            rollback_errors.append(f"{item_id}: {exc}")
    return {"restored_rows": restored, "errors": rollback_errors, "ok": not rollback_errors}


def _postcheck_row(session, row: Mapping[str, Any]) -> dict[str, Any]:
    item_id = int(row["item_id"])
    live = _fetch_inventory_item_by_id(session, item_id)
    if live is None:
        raise UpdateApplyError(f"{row.get('id')}: item no encontrado en postcheck")
    needs_supplier_check = any(
        str(change.get("supplier") or "").strip() and change.get("supplier_price_changed")
        for change in (row.get("changes") or [])
    )
    supplier_map = _supplier_price_map(_fetch_supplier_prices_for_item(session, item_id)) if needs_supplier_check else {}
    checked: list[str] = []
    for change in row.get("changes") or []:
        field = str(change.get("field") or "")
        new_value = change.get("new")
        if change.get("inventory_changed") and not _decimal_equal(live.get(field), new_value):
            raise UpdateApplyError(f"{row.get('id')}: postcheck fallo en {field}")
        supplier = str(change.get("supplier") or "").strip()
        if supplier and change.get("supplier_price_changed"):
            supplier_row = supplier_map.get(supplier) or {}
            if not _decimal_equal(supplier_row.get("price"), new_value):
                raise UpdateApplyError(f"{row.get('id')}: postcheck fallo en supplier_prices[{supplier}]")
        checked.append(field if not supplier else f"{field}+supplier_prices[{supplier}]")
    return {"item_id": item_id, "checked": checked}


def apply_update_preview(session, preview: Mapping[str, Any], settings: Settings | None = None) -> dict[str, Any]:
    if not getattr(session, "user_id", None):
        raise UpdateApplyError("No se puede aplicar Actualizaciones sin usuario autenticado.")
    settings = settings or load_settings()
    process = str(preview.get("process") or "")
    _definition(process)
    if not preview.get("valid"):
        raise UpdateApplyError("El preview contiene errores y no puede aplicarse.")

    ready_rows = [dict(row) for row in (preview.get("rows") or []) if row.get("status") == READY]
    if not ready_rows:
        return {
            "operation_id": "",
            "status": NO_CHANGE,
            "process": process,
            "applied_rows": 0,
            "errors": [],
            "postcheck": [],
        }

    operation_id = new_operation_id("UPDATES")
    now = _now_iso()
    applied: list[dict[str, Any]] = []
    postcheck: list[dict[str, Any]] = []
    error_message = ""

    try:
        for row in ready_rows:
            item_id = int(row["item_id"])
            live = _fetch_inventory_item_by_id(session, item_id)
            if live is None:
                raise UpdateApplyError(f"{row.get('id')}: item no encontrado antes del apply")
            supplier_rows = _fetch_supplier_prices_for_item(session, item_id)
            _validate_live_row(row, live, supplier_rows)

            write_snapshot(
                session,
                OperationSnapshot(
                    operation_id=operation_id,
                    module="Actualizaciones",
                    action=f"apply_{process}",
                    entity_type="inventory_item",
                    entity_id=str(item_id),
                    before_data=_json_safe(
                        {
                            "inventory_items": live,
                            "supplier_prices": supplier_rows,
                            "preview_row": row,
                            "actor": {
                                "user_id": getattr(session, "user_id", None),
                                "email": getattr(session, "email", None),
                                "role": getattr(session, "role", None),
                            },
                        }
                    ),
                    reason="Snapshot before de Actualizaciones antes de escribir Supabase.",
                ),
            )

            inventory_payload: dict[str, Any] = {}
            supplier_updates: list[tuple[str, Any]] = []
            for change in row.get("changes") or []:
                field = str(change.get("field") or "")
                if change.get("inventory_changed"):
                    inventory_payload[field] = change.get("new")
                supplier = str(change.get("supplier") or "").strip()
                if supplier and change.get("supplier_price_changed"):
                    supplier_updates.append((supplier, change.get("new")))

            applied_record = {
                "operation_id": operation_id,
                "item_id": item_id,
                "before_inventory": live,
                "before_supplier_prices": supplier_rows,
                "inventory_fields": sorted(inventory_payload.keys()),
                "supplier_fields": [supplier for supplier, _value in supplier_updates],
            }

            row_tracked_for_rollback = False
            if inventory_payload:
                inventory_payload["updated_at"] = now
                inventory_payload["updated_by"] = getattr(session, "user_id", None)
                _update_inventory_fields(session, item_id, inventory_payload)
                applied.append(applied_record)
                row_tracked_for_rollback = True
            elif supplier_updates:
                applied.append(applied_record)
                row_tracked_for_rollback = True
            for supplier, price in supplier_updates:
                _upsert_supplier_price(session, item_id, supplier, price, operation_id, now)
            if not row_tracked_for_rollback:
                applied.append(applied_record)
            postcheck.append(_postcheck_row(session, row))
    except Exception as exc:
        error_message = str(exc)
        rollback = _rollback_applied_rows(session, applied, now=_now_iso()) if applied else {"restored_rows": 0, "errors": [], "ok": True}
        status = "ROLLED_BACK" if rollback.get("ok") else "FAILED"
        result = {
            "operation_id": operation_id,
            "status": status,
            "process": process,
            "applied_rows": len(applied),
            "errors": [error_message, *rollback.get("errors", [])],
            "rollback": rollback,
            "postcheck": postcheck,
        }
        try:
            write_audit_event(
                session,
                AuditEvent(
                    operation_id=operation_id,
                    module="Actualizaciones",
                    action=f"apply_{process}",
                    status=status,
                    severity="ERROR",
                    entity_type="bulk_update",
                    entity_id=process,
                    before_data={"preview": _json_safe(preview)},
                    after_data=_json_safe(result),
                    message="Actualizacion revertida o fallida tras error.",
                    error_detail=error_message,
                ),
                settings,
            )
        except Exception:
            pass
        raise UpdateApplyError(error_message, result) from exc

    result = {
        "operation_id": operation_id,
        "status": "APPLIED",
        "process": process,
        "applied_rows": len(applied),
        "write_cell_count": sum(len(row.get("changes") or []) for row in ready_rows),
        "updated_fields": sorted({change.get("field") for row in ready_rows for change in (row.get("changes") or []) if change.get("field")}),
        "postcheck": postcheck,
        "errors": [],
    }
    write_audit_event(
        session,
        AuditEvent(
            operation_id=operation_id,
            module="Actualizaciones",
            action=f"apply_{process}",
            status="OK",
            severity="INFO",
            entity_type="bulk_update",
            entity_id=process,
            before_data={"preview": _json_safe(preview)},
            after_data=_json_safe(result),
            message="Actualizacion aplicada desde modulo Actualizaciones. WooCommerce no fue tocado.",
        ),
        settings,
    )
    return result
