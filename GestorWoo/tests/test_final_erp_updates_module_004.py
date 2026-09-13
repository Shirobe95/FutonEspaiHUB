from __future__ import annotations

import inspect
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from futonhub.cloud.services import updates  # noqa: E402


class FakeResponse:
    def __init__(self, data):
        self.data = data


class FakeQuery:
    def __init__(self, client, table_name: str):
        self.client = client
        self.table_name = table_name
        self.filters: list[tuple[str, object]] = []
        self.limit_value: int | None = None
        self.operation = "select"
        self.payload = None

    def select(self, *_args, **_kwargs):
        self.operation = "select"
        return self

    def eq(self, column: str, value):
        self.filters.append((column, value))
        return self

    def limit(self, value: int):
        self.limit_value = int(value)
        return self

    def update(self, payload):
        self.operation = "update"
        self.payload = dict(payload)
        return self

    def upsert(self, payload, **_kwargs):
        self.operation = "upsert"
        self.payload = dict(payload)
        return self

    def delete(self):
        self.operation = "delete"
        return self

    def execute(self):
        if self.operation == "update":
            return self._execute_update()
        if self.operation == "upsert":
            return self._execute_upsert()
        if self.operation == "delete":
            return self._execute_delete()
        return self._execute_select()

    def _rows(self):
        return self.client.tables.setdefault(self.table_name, [])

    def _matches(self, row: dict) -> bool:
        for column, value in self.filters:
            if row.get(column) != value:
                return False
        return True

    def _execute_select(self):
        if self.table_name == "supplier_prices":
            self.client.supplier_select_calls += 1
            if self.client.supplier_select_calls in self.client.fail_supplier_select_on_calls:
                raise RuntimeError("forced supplier_prices select failure")
        rows = [dict(row) for row in self._rows() if self._matches(row)]
        if self.limit_value is not None:
            rows = rows[: self.limit_value]
        return FakeResponse(rows)

    def _execute_update(self):
        matched = []
        for row in self._rows():
            if not self._matches(row):
                continue
            if self.client.fail_inventory_update_item_id == row.get("item_id"):
                raise RuntimeError("forced inventory update failure")
            before_payload = dict(self.payload or {})
            self.client.update_payloads.append((self.table_name, row.get("item_id"), before_payload))
            row.update(before_payload)
            matched.append(dict(row))
        return FakeResponse(matched)

    def _execute_upsert(self):
        payload = dict(self.payload or {})
        self.client.upsert_payloads.append((self.table_name, payload))
        if self.table_name == "supplier_prices":
            supplier = str(payload.get("supplier") or "")
            remaining_failures = self.client.fail_supplier_upsert_once_by_supplier.get(supplier, 0)
            if remaining_failures > 0:
                self.client.fail_supplier_upsert_once_by_supplier[supplier] = remaining_failures - 1
                raise RuntimeError(f"forced supplier_prices upsert failure for {supplier}")
        rows = self._rows()
        if self.table_name == "supplier_prices":
            for row in rows:
                if row.get("item_id") == payload.get("item_id") and row.get("supplier") == payload.get("supplier"):
                    row.update(payload)
                    return FakeResponse([dict(row)])
        rows.append(dict(payload))
        return FakeResponse([payload])

    def _execute_delete(self):
        removed = []
        remaining = []
        for row in self._rows():
            if self._matches(row):
                removed.append(dict(row))
            else:
                remaining.append(row)
        self.client.tables[self.table_name] = remaining
        return FakeResponse(removed)


class FakeClient:
    def __init__(self):
        self.tables = {"inventory_items": [], "supplier_prices": []}
        self.update_payloads: list[tuple[str, object, dict]] = []
        self.upsert_payloads: list[tuple[str, dict]] = []
        self.fail_inventory_update_item_id = None
        self.supplier_select_calls = 0
        self.fail_supplier_select_on_calls: set[int] = set()
        self.fail_supplier_upsert_once_by_supplier: dict[str, int] = {}

    def table(self, name: str):
        return FakeQuery(self, name)


def inventory_row(
    item_id: int,
    code: str,
    *,
    name: str = "Item",
    primary_supplier: str | None = "Ekomat",
    commercial_status: str = "Normal",
    rotation_c=1.0,
    store_stock=2,
    warehouse_stock=3,
    primary_supplier_price=10.0,
    pascal_price=8.0,
):
    source_row = {"primary_supplier": primary_supplier} if primary_supplier else {}
    return {
        "item_id": item_id,
        "name": name,
        "hub_item_code": code,
        "heca_reference": code,
        "woo_sku": code,
        "item_record_type": "simple",
        "base_item_code": "",
        "is_pack": False,
        "commercial_status": commercial_status,
        "rotation_c": rotation_c,
        "store_stock": store_stock,
        "warehouse_stock": warehouse_stock,
        "weighted_average_cost": 99.9,
        "primary_supplier_price": primary_supplier_price,
        "pascal_price": pascal_price,
        "supplier_order_provider": "",
        "source_row": source_row,
        "updated_at": "2026-08-01T00:00:00Z",
    }


def supplier_price(item_id: int, supplier: str, price):
    return {
        "item_id": item_id,
        "supplier": supplier,
        "price": price,
        "currency": "EUR",
        "source": "test",
        "updated_at": "2026-08-01T00:00:00Z",
    }


def fake_session(client: FakeClient, *, email: str = "worker@example.test", user_id: str = "user-1"):
    return SimpleNamespace(user_id=user_id, email=email, role="worker", client=client)


def fake_settings():
    return SimpleNamespace(machine_name="test-machine", sync_role="worker")


def write_futon_espai_stock_workbook(path: Path, rows: list[dict[str, object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet2"
    ws.append(["Codigo", "Denominacion", "Cantidad", "Coste", "Valoracion", "1-NADAL FORW", "2-FUTONESPAI", "Ignored"])
    for row in rows:
        ws.append(
            [
                row.get("code"),
                row.get("name", "Item"),
                row.get("total"),
                row.get("cost", 99),
                row.get("valuation", 999),
                row.get("other_warehouse", 123),
                row.get("store"),
                row.get("ignored", "ignored"),
            ]
        )
        number_format = row.get("number_format")
        if number_format:
            ws.cell(row=ws.max_row, column=1).number_format = str(number_format)
    wb.create_sheet("Sheet1")
    wb.save(path)
    wb.close()


class UpdatesModule004Tests(unittest.TestCase):
    def test_updates_access_policy_is_exact_and_limited_to_configured_users(self) -> None:
        all_processes = (
            updates.PROCESS_ROTATION_C,
            updates.PROCESS_STOCK,
            updates.PROCESS_SUPPLIER_PRICES,
        )

        self.assertEqual(updates.allowed_updates_processes("andyshb95@gmail.com"), all_processes)
        self.assertEqual(updates.allowed_updates_processes("ANDYSHB95@GMAIL.COM"), all_processes)
        self.assertEqual(updates.allowed_updates_processes("futonhub1@gmail.com"), (updates.PROCESS_STOCK,))
        self.assertEqual(updates.allowed_updates_processes("FUTONHUB1@GMAIL.COM"), (updates.PROCESS_STOCK,))
        self.assertEqual(updates.allowed_updates_processes("xfutonhub1@gmail.com"), all_processes)
        self.assertEqual(updates.allowed_updates_processes("worker@example.test"), all_processes)

    def test_futon_espai_can_preview_and_apply_stock(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", store_stock=2, warehouse_stock=3)]
        session = fake_session(client, email="futonhub1@gmail.com")

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [{"ID": "0725002", "Stock Tienda": 7, "Stock Warehouse": 8}],
        )

        self.assertTrue(preview["apply_enabled"])
        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        self.assertEqual(result["applied_rows"], 1)
        self.assertEqual(client.tables["inventory_items"][0]["store_stock"], 7.0)
        self.assertEqual(client.tables["inventory_items"][0]["warehouse_stock"], 8.0)

    def test_futon_espai_rotation_and_supplier_prices_are_access_denied_before_writes(self) -> None:
        for process in (updates.PROCESS_ROTATION_C, updates.PROCESS_SUPPLIER_PRICES):
            with self.subTest(process=process):
                client = FakeClient()
                client.tables["inventory_items"] = [inventory_row(725002, "0725002", rotation_c=1.0)]
                session = fake_session(client, email="futonhub1@gmail.com")

                with self.assertRaises(updates.UpdateApplyError) as preview_ctx:
                    updates.preview_updates(session, process, [{"ID": "0725002", "Rotación C": 1.5}])

                self.assertEqual(preview_ctx.exception.result["status"], updates.ACCESS_DENIED)
                stale_ready_preview = {
                    "process": process,
                    "valid": True,
                    "global_blocker": False,
                    "rows": [
                        {
                            "process": process,
                            "status": updates.READY,
                            "id": "0725002",
                            "item_id": 725002,
                            "changes": [{"field": "rotation_c", "new": 1.5, "inventory_changed": True}],
                        }
                    ],
                }
                with patch.object(updates, "write_snapshot") as snapshots, patch.object(updates, "write_audit_event") as audits:
                    with self.assertRaises(updates.UpdateApplyError) as apply_ctx:
                        updates.apply_update_preview(session, stale_ready_preview, settings=fake_settings())

                self.assertEqual(apply_ctx.exception.result["status"], updates.ACCESS_DENIED)
                snapshots.assert_not_called()
                audits.assert_not_called()
                self.assertEqual(client.update_payloads, [])
                self.assertEqual(client.upsert_payloads, [])

    def test_owner_keeps_full_updates_access(self) -> None:
        self.assertTrue(updates.can_use_updates_process("andyshb95@gmail.com", updates.PROCESS_STOCK))
        self.assertTrue(updates.can_use_updates_process("andyshb95@gmail.com", updates.PROCESS_ROTATION_C))
        self.assertTrue(updates.can_use_updates_process("andyshb95@gmail.com", updates.PROCESS_SUPPLIER_PRICES))

    def test_template_generation_and_excel_parse_preserve_leading_zeros(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            template_path = Path(tmp) / "Plantilla_RotacionC.xlsx"
            updates.write_update_template(updates.PROCESS_ROTATION_C, template_path)
            wb = load_workbook(template_path)
            ws = wb.active
            self.assertEqual([ws.cell(1, 1).value, ws.cell(1, 2).value], ["ID", "Rotación C"])
            self.assertIsNone(ws.cell(2, 1).value)
            wb.close()

            data_path = Path(tmp) / "datos.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["ID", "Rotación C"])
            ws.append(["0725002", "1,25"])
            ws.append([725003, 2])
            ws.cell(row=3, column=1).number_format = "0000000"
            wb.save(data_path)

            rows = updates.read_updates_workbook(data_path, updates.PROCESS_ROTATION_C)
            self.assertEqual(rows[0]["ID"], "0725002")
            self.assertEqual(rows[1]["ID"], "0725003")

            bad_path = Path(tmp) / "headers_invalidos.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append([" ID", "Rotación C"])
            wb.save(bad_path)
            with self.assertRaises(updates.UpdatesValidationError):
                updates.read_updates_workbook(bad_path, updates.PROCESS_ROTATION_C)

    def test_rotation_preview_no_change_invalid_duplicate_and_apply_mapping(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002", rotation_c=1.0),
            inventory_row(725003, "0725003", rotation_c=2.0),
        ]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [
                {"ID": "0725002", "Rotación C": "1,50"},
                {"ID": "0725003", "Rotación C": 2},
            ],
        )
        self.assertTrue(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.READY: 1, updates.NO_CHANGE: 1})

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        self.assertEqual(client.tables["inventory_items"][0]["rotation_c"], 1.5)
        payload_fields = {key for _table, _item_id, payload in client.update_payloads for key in payload}
        self.assertIn("rotation_c", payload_fields)
        self.assertNotIn("weighted_average_cost", payload_fields)

        blocked = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [
                {"ID": "9999999", "Rotación C": 1},
                {"ID": "0725002", "Rotación C": "mal"},
                {"ID": "0725003", "Rotación C": 1},
                {"ID": "0725003", "Rotación C": 1},
            ],
        )
        self.assertFalse(blocked["apply_enabled"])
        self.assertEqual(blocked["status_counts"][updates.INVALID_ID], 1)
        self.assertEqual(blocked["status_counts"][updates.INVALID_VALUE], 1)
        self.assertEqual(blocked["status_counts"][updates.DUPLICATE_ID], 2)

    def test_stock_preview_apply_descatalogado_blank_and_no_weighted_cost_write(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002", store_stock=2, warehouse_stock=3),
            inventory_row(725004, "0725004", commercial_status="Descatalogado"),
        ]
        session = fake_session(client)

        blocked = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"ID": "0725004", "Stock Tienda": 1, "Stock Warehouse": 1},
                {"ID": "0725002", "Stock Tienda": "", "Stock Warehouse": 4},
            ],
        )
        self.assertFalse(blocked["apply_enabled"])
        self.assertEqual(blocked["status_counts"][updates.DESCATALOGADO_NO_ACTUALIZABLE], 1)
        self.assertEqual(blocked["status_counts"][updates.INVALID_VALUE], 1)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [{"ID": "0725002", "Stock Tienda": 7, "Stock Warehouse": "9"}],
        )
        self.assertEqual(preview["rows"][0]["stock_total_current"], 5.0)
        self.assertEqual(preview["rows"][0]["stock_total_new"], 16.0)
        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            updates.apply_update_preview(session, preview, settings=fake_settings())
        row = client.tables["inventory_items"][0]
        self.assertEqual(row["store_stock"], 7.0)
        self.assertEqual(row["warehouse_stock"], 9.0)
        self.assertEqual(row["weighted_average_cost"], 99.9)
        payload_fields = {key for _table, _item_id, payload in client.update_payloads for key in payload}
        self.assertNotIn("weighted_average_cost", payload_fields)

    def test_futon_espai_stock_excel_derives_warehouse_and_preserves_leading_zero(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Stock Almacenes.xlsx"
            write_futon_espai_stock_workbook(
                path,
                [
                    {"code": "0725002", "total": 13, "store": 4, "ignored": 500},
                    {"code": 725003, "number_format": "0000000", "total": 10, "store": 10},
                    {"code": "0725004", "total": 0, "store": 0},
                ],
            )

            rows = updates.read_updates_workbook(path, updates.PROCESS_STOCK)

        self.assertEqual([row["ID"] for row in rows], ["0725002", "0725003", "0725004"])
        self.assertEqual(rows[0]["Stock Tienda"], 4)
        self.assertEqual(rows[0]["Stock Warehouse"], 9.0)
        self.assertEqual(rows[0]["_stock_total_excel"], 13)
        self.assertEqual(rows[1]["Stock Warehouse"], 0.0)
        self.assertEqual(rows[2]["Stock Warehouse"], 0.0)
        self.assertEqual(rows[0]["_stock_import_format"], updates.STOCK_FORMAT_FUTON_ESPAI)

    def test_futon_espai_stock_preview_uses_excel_total_and_writes_only_stock_fields(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Stock Almacenes.xlsx"
            write_futon_espai_stock_workbook(path, [{"code": "0725002", "total": 13, "store": 4}])
            rows = updates.read_updates_workbook(path, updates.PROCESS_STOCK)

        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", store_stock=1, warehouse_stock=2)]
        session = fake_session(client)
        preview = updates.preview_updates(session, updates.PROCESS_STOCK, rows)

        self.assertTrue(preview["apply_enabled"])
        row = preview["rows"][0]
        self.assertEqual(row["store_stock_new"], 4.0)
        self.assertEqual(row["warehouse_stock_new"], 9.0)
        self.assertEqual(row["stock_total_new"], 13.0)
        self.assertEqual(row["stock_total_excel"], 13.0)

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        self.assertEqual(result["updated_fields"], ["store_stock", "warehouse_stock"])
        self.assertEqual(client.tables["inventory_items"][0]["store_stock"], 4.0)
        self.assertEqual(client.tables["inventory_items"][0]["warehouse_stock"], 9.0)
        self.assertEqual(client.tables["inventory_items"][0]["weighted_average_cost"], 99.9)
        payload_fields = {key for _table, _item_id, payload in client.update_payloads for key in payload}
        self.assertEqual(payload_fields, {"store_stock", "warehouse_stock", "updated_at", "updated_by"})

    def test_futon_espai_stock_signed_values_are_valid(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(201001, "0201001", store_stock=0, warehouse_stock=0),
            inventory_row(201011, "0201011", store_stock=0, warehouse_stock=0),
            inventory_row(808001, "0808001", store_stock=0, warehouse_stock=0),
            inventory_row(817003, "0817003", store_stock=0, warehouse_stock=0),
            inventory_row(725004, "0725004", store_stock=1, warehouse_stock=1),
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"_row_number": 2, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0201001", "_stock_total_excel": 29, "Stock Tienda": -1},
                {"_row_number": 3, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0201011", "_stock_total_excel": -197, "Stock Tienda": 46},
                {"_row_number": 4, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0808001", "_stock_total_excel": -14, "Stock Tienda": 3},
                {"_row_number": 5, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0817003", "_stock_total_excel": 1, "Stock Tienda": 2},
                {"_row_number": 6, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725004", "_stock_total_excel": 0, "Stock Tienda": 0},
            ],
        )

        self.assertTrue(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.READY: 5})
        by_id = {row["id"]: row for row in preview["rows"]}
        self.assertEqual(by_id["0201001"]["warehouse_stock_new"], 30.0)
        self.assertEqual(by_id["0201011"]["warehouse_stock_new"], -243.0)
        self.assertEqual(by_id["0808001"]["warehouse_stock_new"], -17.0)
        self.assertEqual(by_id["0817003"]["warehouse_stock_new"], -1.0)
        self.assertEqual(by_id["0725004"]["warehouse_stock_new"], 0.0)

    def test_futon_espai_stock_empty_non_numeric_and_decimal_values_are_blocked(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002"),
            inventory_row(725003, "0725003"),
            inventory_row(725004, "0725004"),
            inventory_row(725005, "0725005"),
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"_row_number": 2, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725002", "_stock_total_excel": "", "Stock Tienda": 1},
                {"_row_number": 3, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725003", "_stock_total_excel": 10, "Stock Tienda": ""},
                {"_row_number": 4, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725004", "_stock_total_excel": "mal", "Stock Tienda": 0},
                {"_row_number": 5, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725005", "_stock_total_excel": 10, "Stock Tienda": "1.5"},
            ],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.INVALID_VALUE: 4})
        reasons = " | ".join(row["reason"] for row in preview["rows"])
        self.assertIn("Total Excel vacio", reasons)
        self.assertIn("Stock Tienda vacio", reasons)
        self.assertIn("Total Excel no es numerico", reasons)
        self.assertIn("Stock Tienda debe ser un entero", reasons)

    def test_mixed_stock_row_errors_keep_apply_enabled_for_ready_rows_only(self) -> None:
        client = FakeClient()
        ready_codes = [f"09010{index:02d}" for index in range(10)]
        no_change_codes = [f"09020{index:02d}" for index in range(3)]
        descatalogado_codes = [f"09030{index:02d}" for index in range(2)]
        invalid_value_code = "0904000"
        client.tables["inventory_items"] = [
            *[
                inventory_row(901000 + index, code, store_stock=0, warehouse_stock=0)
                for index, code in enumerate(ready_codes)
            ],
            *[
                inventory_row(902000 + index, code, store_stock=3, warehouse_stock=7)
                for index, code in enumerate(no_change_codes)
            ],
            *[
                inventory_row(903000 + index, code, commercial_status="Descatalogado")
                for index, code in enumerate(descatalogado_codes)
            ],
            inventory_row(904000, invalid_value_code),
        ]
        session = fake_session(client)

        rows = [
            *[
                {
                    "_row_number": index + 2,
                    "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                    "ID": code,
                    "_stock_total_excel": 15,
                    "Stock Tienda": 4,
                }
                for index, code in enumerate(ready_codes)
            ],
            *[
                {
                    "_row_number": index + 20,
                    "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                    "ID": code,
                    "_stock_total_excel": 10,
                    "Stock Tienda": 3,
                }
                for index, code in enumerate(no_change_codes)
            ],
            *[
                {
                    "_row_number": index + 30,
                    "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                    "ID": code,
                    "_stock_total_excel": 10,
                    "Stock Tienda": 3,
                }
                for index, code in enumerate(descatalogado_codes)
            ],
            {
                "_row_number": 40,
                "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                "ID": "0999990",
                "_stock_total_excel": 1,
                "Stock Tienda": 0,
            },
            {
                "_row_number": 41,
                "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                "ID": "0999991",
                "_stock_total_excel": 1,
                "Stock Tienda": 0,
            },
            {
                "_row_number": 42,
                "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                "ID": invalid_value_code,
                "_stock_total_excel": 1,
                "Stock Tienda": "1.5",
            },
        ]

        preview = updates.preview_updates(session, updates.PROCESS_STOCK, rows)

        self.assertTrue(preview["valid"])
        self.assertTrue(preview["apply_enabled"])
        self.assertEqual(preview["ready_count"], 10)
        self.assertEqual(preview["no_change_count"], 3)
        self.assertEqual(preview["row_error_count"], 5)
        self.assertEqual(preview["excluded_count"], 8)
        self.assertEqual(
            preview["status_counts"],
            {
                updates.READY: 10,
                updates.NO_CHANGE: 3,
                updates.DESCATALOGADO_NO_ACTUALIZABLE: 2,
                updates.INVALID_ID: 2,
                updates.INVALID_VALUE: 1,
            },
        )

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        self.assertEqual(result["applied_rows"], 10)
        self.assertEqual({item_id for _table, item_id, _payload in client.update_payloads}, {901000 + index for index in range(10)})

    def test_global_blocker_still_blocks_apply_even_with_ready_rows(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", store_stock=0, warehouse_stock=0)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {
                    "_row_number": 2,
                    "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI,
                    "ID": "0725002",
                    "_stock_total_excel": 10,
                    "Stock Tienda": 3,
                }
            ],
        )
        preview["global_blocker"] = True
        preview["valid"] = False

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError):
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(client.update_payloads, [])

    def test_futon_espai_stock_apply_can_write_signed_store_and_warehouse(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(201001, "0201001", store_stock=0, warehouse_stock=0),
            inventory_row(201011, "0201011", store_stock=0, warehouse_stock=0),
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"_row_number": 2, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0201001", "_stock_total_excel": 29, "Stock Tienda": -1},
                {"_row_number": 3, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0201011", "_stock_total_excel": -197, "Stock Tienda": 46},
            ],
        )

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        by_item = {row["item_id"]: row for row in client.tables["inventory_items"]}
        self.assertEqual(by_item[201001]["store_stock"], -1.0)
        self.assertEqual(by_item[201001]["warehouse_stock"], 30.0)
        self.assertEqual(by_item[201011]["store_stock"], 46.0)
        self.assertEqual(by_item[201011]["warehouse_stock"], -243.0)
        payloads = {item_id: payload for _table, item_id, payload in client.update_payloads}
        self.assertEqual(payloads[201001]["store_stock"], -1.0)
        self.assertEqual(payloads[201001]["warehouse_stock"], 30.0)
        self.assertEqual(payloads[201011]["store_stock"], 46.0)
        self.assertEqual(payloads[201011]["warehouse_stock"], -243.0)

    def test_stock_resolved_duplicate_item_ids_are_excluded_before_apply(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(201004, "0201004", store_stock=0, warehouse_stock=9)]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"_row_number": 8, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0201004", "_stock_total_excel": 18, "Stock Tienda": 0},
                {"_row_number": 265, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "201004", "_stock_total_excel": 0, "Stock Tienda": 0},
            ],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["ready_count"], 0)
        self.assertEqual(preview["write_cell_count"], 0)
        self.assertEqual(preview["status_counts"], {updates.DUPLICATE_ID: 2})
        self.assertEqual({row["item_id"] for row in preview["rows"]}, {201004})
        self.assertTrue(all("item_id=201004" in row["reason"] for row in preview["rows"]))

    def test_stale_ready_preview_with_duplicate_item_ids_blocks_without_writes(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(201004, "0201004", store_stock=0, warehouse_stock=9)]
        session = fake_session(client)
        stale_preview = {
            "process": updates.PROCESS_STOCK,
            "valid": True,
            "global_blocker": False,
            "rows": [
                {
                    "process": updates.PROCESS_STOCK,
                    "status": updates.READY,
                    "id": "0201004",
                    "requested_id": "0201004",
                    "item_id": 201004,
                    "changes": [
                        {
                            "field": "warehouse_stock",
                            "current": 9,
                            "new": 18.0,
                            "expected_current": 9,
                            "inventory_changed": True,
                        }
                    ],
                },
                {
                    "process": updates.PROCESS_STOCK,
                    "status": updates.READY,
                    "id": "0201004",
                    "requested_id": "201004",
                    "item_id": 201004,
                    "changes": [
                        {
                            "field": "warehouse_stock",
                            "current": 9,
                            "new": 0.0,
                            "expected_current": 9,
                            "inventory_changed": True,
                        }
                    ],
                },
            ],
        }

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError) as ctx:
                updates.apply_update_preview(session, stale_preview, settings=fake_settings())

        self.assertEqual(ctx.exception.result["status"], "BLOCKED_DUPLICATE_READY_ITEM")
        self.assertEqual(ctx.exception.result["applied_rows"], 0)
        self.assertEqual(client.tables["inventory_items"][0]["warehouse_stock"], 9)
        self.assertEqual(client.update_payloads, [])

    def test_stock_apply_two_fields_same_row_does_not_self_drift(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", store_stock=1, warehouse_stock=10)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [{"ID": "0725002", "Stock Tienda": 2, "Stock Warehouse": 9}],
        )

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        row = client.tables["inventory_items"][0]
        self.assertEqual(row["store_stock"], 2.0)
        self.assertEqual(row["warehouse_stock"], 9.0)
        self.assertEqual(len(client.update_payloads), 1)
        payload = client.update_payloads[0][2]
        self.assertEqual(payload["store_stock"], 2.0)
        self.assertEqual(payload["warehouse_stock"], 9.0)

    def test_stock_apply_real_external_drift_still_blocks_and_rolls_back(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002", store_stock=1, warehouse_stock=10),
            inventory_row(725003, "0725003", store_stock=4, warehouse_stock=10),
        ]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"ID": "0725002", "Stock Tienda": 2, "Stock Warehouse": 9},
                {"ID": "0725003", "Stock Tienda": 4, "Stock Warehouse": 12},
            ],
        )
        client.tables["inventory_items"][1]["warehouse_stock"] = 11

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError) as ctx:
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(ctx.exception.result["status"], "ROLLED_BACK")
        by_item = {row["item_id"]: row for row in client.tables["inventory_items"]}
        self.assertEqual(by_item[725002]["store_stock"], 1)
        self.assertEqual(by_item[725002]["warehouse_stock"], 10)
        self.assertEqual(by_item[725003]["warehouse_stock"], 11)

    def test_stock_apply_numeric_equivalent_values_do_not_trigger_drift(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", store_stock=0, warehouse_stock=1)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [{"ID": "0725002", "Stock Tienda": 2, "Stock Warehouse": 3}],
        )
        client.tables["inventory_items"][0]["store_stock"] = Decimal("0.0")
        client.tables["inventory_items"][0]["warehouse_stock"] = Decimal("1.0")

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        row = client.tables["inventory_items"][0]
        self.assertEqual(row["store_stock"], 2.0)
        self.assertEqual(row["warehouse_stock"], 3.0)

    def test_futon_espai_stock_duplicates_invalid_id_and_descatalogado_are_preserved(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002"),
            inventory_row(725004, "0725004", commercial_status="Descatalogado"),
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"_row_number": 2, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725002", "_stock_total_excel": 13, "Stock Tienda": 4},
                {"_row_number": 3, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725002", "_stock_total_excel": 14, "Stock Tienda": 5},
                {"_row_number": 4, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "9999999", "_stock_total_excel": 1, "Stock Tienda": 0},
                {"_row_number": 5, "_stock_import_format": updates.STOCK_FORMAT_FUTON_ESPAI, "ID": "0725004", "_stock_total_excel": 1, "Stock Tienda": 0},
            ],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"][updates.DUPLICATE_ID], 2)
        self.assertEqual(preview["status_counts"][updates.INVALID_ID], 1)
        self.assertEqual(preview["status_counts"][updates.DESCATALOGADO_NO_ACTUALIZABLE], 1)

    def test_futon_espai_stock_formula_without_cached_value_is_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Stock Almacenes.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet2"
            ws.append(["Codigo", "Denominacion", "Cantidad", "Coste", "Valoracion", "1-NADAL FORW", "2-FUTONESPAI"])
            ws.append(["0725002", "Item", "=SUM(F2:Q2)", 0, 0, 9, 4])
            wb.save(path)
            wb.close()

            with self.assertRaises(updates.UpdatesValidationError) as ctx:
                updates.read_updates_workbook(path, updates.PROCESS_STOCK)

        self.assertIn("formula sin valor calculado", str(ctx.exception))
        self.assertIn("fila 2", str(ctx.exception))

    def test_stock_workbook_with_unknown_signature_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "stock_unknown.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Codigo", "Nombre", "Cantidad", "", "", "", "Tienda"])
            ws.append(["0725002", "Item", 13, "", "", "", 4])
            wb.save(path)
            wb.close()

            with self.assertRaises(updates.UpdatesValidationError) as ctx:
                updates.read_updates_workbook(path, updates.PROCESS_STOCK)

        self.assertIn("Stock Almacenes de Futon Espai", str(ctx.exception))

    def test_legacy_stock_workbook_format_is_still_supported(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy_stock.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["ID", "Stock Tienda", "Stock Warehouse"])
            ws.append(["0725002", 7, 8])
            wb.save(path)
            wb.close()

            rows = updates.read_updates_workbook(path, updates.PROCESS_STOCK)

        self.assertEqual(rows, [{"_row_number": 2, "ID": "0725002", "Stock Tienda": 7, "Stock Warehouse": 8}])

    def test_supplier_prices_primary_resolution_pascal_blank_and_sync(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002", primary_supplier="Ekomat", primary_supplier_price=10, pascal_price=8),
            inventory_row(302009, "0302009", primary_supplier="Cipta", primary_supplier_price=20, pascal_price=18),
            inventory_row(201001, "0201001", primary_supplier="Hemei", primary_supplier_price=30, pascal_price=28),
            inventory_row(999001, "0999001", primary_supplier=None, primary_supplier_price=11, pascal_price=9),
        ]
        client.tables["supplier_prices"] = [
            supplier_price(725002, "Ekomat", 10),
            supplier_price(725002, "Pascal", 8),
            supplier_price(302009, "Cipta", 20),
            supplier_price(201001, "Hemei", 30),
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [
                {"ID": "0725002", "Precio Principal": "10,50", "Pascal": ""},
                {"ID": "0302009", "Precio Principal": "21.25", "Pascal": ""},
                {"ID": "0201001", "Precio Principal": "", "Pascal": "29,75"},
            ],
        )
        self.assertEqual(preview["status_counts"], {updates.READY: 3})
        self.assertEqual([row["primary_supplier"] for row in preview["rows"]], ["Ekomat", "Cipta", "Hemei"])

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(result["status"], "APPLIED")
        by_item = {row["item_id"]: row for row in client.tables["inventory_items"]}
        self.assertEqual(by_item[725002]["primary_supplier_price"], 10.5)
        self.assertEqual(by_item[302009]["primary_supplier_price"], 21.25)
        self.assertEqual(by_item[201001]["pascal_price"], 29.75)
        supplier_lookup = {(row["item_id"], row["supplier"]): row["price"] for row in client.tables["supplier_prices"]}
        self.assertEqual(supplier_lookup[(725002, "Ekomat")], 10.5)
        self.assertEqual(supplier_lookup[(302009, "Cipta")], 21.25)
        self.assertEqual(supplier_lookup[(201001, "Pascal")], 29.75)

        blocked = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [
                {"ID": "0999001", "Precio Principal": "12", "Pascal": ""},
                {"ID": "0725002", "Precio Principal": "", "Pascal": ""},
            ],
        )
        self.assertFalse(blocked["apply_enabled"])
        self.assertEqual(blocked["status_counts"][updates.PRIMARY_SUPPLIER_UNRESOLVED], 1)
        self.assertEqual(blocked["status_counts"][updates.NO_CHANGE], 1)

    def test_six_digit_numeric_id_can_resolve_unique_zero_prefixed_physical_item(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002"),
            inventory_row(758087, "0758087"),
            inventory_row(1020007, "1020007"),
        ]
        session = fake_session(client)

        valid = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [{"ID": "0725002", "Rotación C": 1.5}],
        )

        self.assertTrue(valid["apply_enabled"])
        self.assertEqual(valid["rows"][0]["item_id"], 725002)
        self.assertEqual(valid["rows"][0]["id"], "0725002")

        without_zero = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [{"ID": "725002", "Rotación C": 1.5}],
        )

        self.assertTrue(without_zero["apply_enabled"])
        self.assertEqual(without_zero["rows"][0]["item_id"], 725002)
        self.assertEqual(without_zero["rows"][0]["id"], "0725002")
        self.assertEqual(without_zero["rows"][0]["requested_id"], "725002")

        invalid = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [
                {"ID": "0758007", "Rotación C": 1},
                {"ID": "1020011", "Rotación C": 1},
            ],
        )

        self.assertFalse(invalid["apply_enabled"])
        self.assertEqual(invalid["status_counts"][updates.INVALID_ID], 2)

    def test_alias_with_same_heca_reference_does_not_make_unique_physical_item_ambiguous(self) -> None:
        client = FakeClient()
        alias = inventory_row(724012, "0724011A")
        alias.update(
            {
                "heca_reference": "0724011",
                "item_record_type": "alias",
                "base_item_code": "0724011",
            }
        )
        client.tables["inventory_items"] = [
            inventory_row(724011, "0724011"),
            alias,
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [{"ID": "724011", "Rotación C": 1.5}],
        )

        self.assertTrue(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.READY: 1})
        self.assertEqual(preview["rows"][0]["item_id"], 724011)
        self.assertEqual(preview["rows"][0]["id"], "0724011")

    def test_multiple_physical_exact_candidates_still_block_as_ambiguous(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(724011, "0724011"),
            inventory_row(724099, "0724011"),
        ]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [{"ID": "0724011", "Rotación C": 1.5}],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.AMBIGUOUS_IDENTITY: 1})
        self.assertIn("ambigua", preview["rows"][0]["reason"])

    def test_descatalogado_identity_resolves_before_policy_even_if_not_simple(self) -> None:
        client = FakeClient()
        discontinued = inventory_row(725001, "0725001", commercial_status="Descatalogado")
        discontinued["item_record_type"] = "historical_duplicate"
        client.tables["inventory_items"] = [discontinued]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [{"ID": "0725001", "Stock Tienda": 1, "Stock Warehouse": 1}],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.DESCATALOGADO_NO_ACTUALIZABLE: 1})
        self.assertEqual(preview["rows"][0]["item_id"], 725001)
        self.assertEqual(preview["rows"][0]["id"], "0725001")

    def test_zero_prefixed_resolution_does_not_bypass_descatalogado_block(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725001, "0725001", commercial_status="Descatalogado")]
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [{"ID": "725001", "Stock Tienda": 1, "Stock Warehouse": 1}],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["status_counts"], {updates.DESCATALOGADO_NO_ACTUALIZABLE: 1})
        self.assertEqual(preview["rows"][0]["id"], "0725001")

    def test_supplier_price_upsert_failure_rolls_back_current_inventory_update(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", primary_supplier_price=10)]
        client.tables["supplier_prices"] = [supplier_price(725002, "Ekomat", 10)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [{"ID": "0725002", "Precio Principal": 11, "Pascal": ""}],
        )
        client.fail_supplier_upsert_once_by_supplier = {"Ekomat": 1}

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError) as ctx:
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(ctx.exception.result["status"], "ROLLED_BACK")
        self.assertEqual(client.tables["inventory_items"][0]["primary_supplier_price"], 10)
        supplier_lookup = {(row["item_id"], row["supplier"]): row["price"] for row in client.tables["supplier_prices"]}
        self.assertEqual(supplier_lookup[(725002, "Ekomat")], 10)

    def test_pascal_failure_rolls_back_inventory_primary_supplier_and_pascal(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", primary_supplier_price=10, pascal_price=8)]
        client.tables["supplier_prices"] = [
            supplier_price(725002, "Ekomat", 10),
            supplier_price(725002, "Pascal", 8),
        ]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [{"ID": "0725002", "Precio Principal": 11, "Pascal": 9}],
        )
        client.fail_supplier_upsert_once_by_supplier = {"Pascal": 1}

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError) as ctx:
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(ctx.exception.result["status"], "ROLLED_BACK")
        item = client.tables["inventory_items"][0]
        self.assertEqual(item["primary_supplier_price"], 10)
        self.assertEqual(item["pascal_price"], 8)
        supplier_lookup = {(row["item_id"], row["supplier"]): row["price"] for row in client.tables["supplier_prices"]}
        self.assertEqual(supplier_lookup[(725002, "Ekomat")], 10)
        self.assertEqual(supplier_lookup[(725002, "Pascal")], 8)

    def test_supplier_prices_read_failure_blocks_preview(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", primary_supplier_price=10)]
        client.fail_supplier_select_on_calls = {1}
        session = fake_session(client)

        preview = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [{"ID": "0725002", "Precio Principal": 11, "Pascal": ""}],
        )

        self.assertFalse(preview["apply_enabled"])
        self.assertEqual(preview["rows"][0]["status"], updates.SUPPLIER_PRICES_READ_ERROR)
        self.assertIn(updates.SUPPLIER_PRICES_READ_ERROR, updates.BLOCKING_STATUSES)

    def test_supplier_prices_read_failure_before_apply_writes_zero_rows(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", primary_supplier_price=10)]
        client.tables["supplier_prices"] = [supplier_price(725002, "Ekomat", 10)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [{"ID": "0725002", "Precio Principal": 11, "Pascal": ""}],
        )
        client.fail_supplier_select_on_calls = {client.supplier_select_calls + 1}

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError):
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(client.tables["inventory_items"][0]["primary_supplier_price"], 10)
        self.assertEqual(client.update_payloads, [])
        self.assertEqual(client.upsert_payloads, [])

    def test_supplier_prices_read_failure_during_postcheck_rolls_back(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", primary_supplier_price=10)]
        client.tables["supplier_prices"] = [supplier_price(725002, "Ekomat", 10)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_SUPPLIER_PRICES,
            [{"ID": "0725002", "Precio Principal": 11, "Pascal": ""}],
        )
        client.fail_supplier_select_on_calls = {client.supplier_select_calls + 2}

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError) as ctx:
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(ctx.exception.result["status"], "ROLLED_BACK")
        self.assertEqual(client.tables["inventory_items"][0]["primary_supplier_price"], 10)
        supplier_lookup = {(row["item_id"], row["supplier"]): row["price"] for row in client.tables["supplier_prices"]}
        self.assertEqual(supplier_lookup[(725002, "Ekomat")], 10)

    def test_drift_before_apply_blocks_without_write(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", rotation_c=1.0)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [{"ID": "0725002", "Rotación C": 1.5}],
        )
        client.tables["inventory_items"][0]["rotation_c"] = 1.2

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError):
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(client.tables["inventory_items"][0]["rotation_c"], 1.2)
        self.assertEqual(client.update_payloads, [])

    def test_partial_failure_rolls_back_already_applied_rows(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [
            inventory_row(725002, "0725002", store_stock=2, warehouse_stock=3),
            inventory_row(725003, "0725003", store_stock=4, warehouse_stock=5),
        ]
        client.fail_inventory_update_item_id = 725003
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_STOCK,
            [
                {"ID": "0725002", "Stock Tienda": 7, "Stock Warehouse": 8},
                {"ID": "0725003", "Stock Tienda": 9, "Stock Warehouse": 10},
            ],
        )

        with patch.object(updates, "write_snapshot"), patch.object(updates, "write_audit_event"):
            with self.assertRaises(updates.UpdateApplyError) as ctx:
                updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(ctx.exception.result["status"], "ROLLED_BACK")
        by_item = {row["item_id"]: row for row in client.tables["inventory_items"]}
        self.assertEqual(by_item[725002]["store_stock"], 2)
        self.assertEqual(by_item[725002]["warehouse_stock"], 3)
        self.assertEqual(by_item[725003]["store_stock"], 4)
        self.assertEqual(by_item[725003]["warehouse_stock"], 5)

    def test_snapshot_audit_actor_and_postcheck_are_recorded(self) -> None:
        client = FakeClient()
        client.tables["inventory_items"] = [inventory_row(725002, "0725002", rotation_c=1.0)]
        session = fake_session(client)
        preview = updates.preview_updates(
            session,
            updates.PROCESS_ROTATION_C,
            [{"ID": "0725002", "Rotación C": 1.5}],
        )

        with patch.object(updates, "write_snapshot") as snapshots, patch.object(updates, "write_audit_event") as audits:
            result = updates.apply_update_preview(session, preview, settings=fake_settings())

        self.assertEqual(len(result["postcheck"]), 1)
        snapshot_arg = snapshots.call_args.args[1]
        self.assertEqual(snapshot_arg.before_data["actor"]["user_id"], "user-1")
        audit_arg = audits.call_args.args[1]
        self.assertEqual(audit_arg.status, "OK")
        self.assertEqual(audit_arg.after_data["applied_rows"], 1)


class UpdatesNavigationContractTests(unittest.TestCase):
    def test_updates_navigation_entry_and_builder_exist(self) -> None:
        from futonhub.ui.erp.prototype import FutonHubErpPrototype, NAV_ITEMS

        self.assertIn(("actualizaciones", "Actualizaciones", "Operaciones"), [(item.key, item.label, item.group) for item in NAV_ITEMS])
        self.assertTrue(hasattr(FutonHubErpPrototype, "_build_updates"))

    def test_updates_apply_runs_in_background_and_blocks_double_apply(self) -> None:
        from futonhub.ui.erp.prototype import FutonHubErpPrototype

        source = inspect.getsource(FutonHubErpPrototype._build_updates)
        self.assertIn("_updates_process_options_for_user", source)
        self.assertIn('"loading": False', source)
        self.assertIn('if state.get("applying") or state.get("loading"):', source)
        self.assertIn('self._show_working_overlay("Aplicando actualizacion"', source)
        self.assertIn("threading.Thread(target=worker, daemon=True).start()", source)
        self.assertIn("self.after(0, callback)", source)
        self.assertIn("preview_updates_from_excel(session, process, source_path)", source)

    def test_updates_visible_process_options_follow_access_policy(self) -> None:
        from futonhub.ui.erp.prototype import _updates_process_options_for_user

        andy_processes = tuple(process for process, _label in _updates_process_options_for_user("andyshb95@gmail.com"))
        espai_processes = tuple(process for process, _label in _updates_process_options_for_user("futonhub1@gmail.com"))
        unknown_processes = tuple(process for process, _label in _updates_process_options_for_user("worker@example.test"))

        expected_all = (
            updates.PROCESS_ROTATION_C,
            updates.PROCESS_STOCK,
            updates.PROCESS_SUPPLIER_PRICES,
        )
        self.assertEqual(andy_processes, expected_all)
        self.assertEqual(espai_processes, (updates.PROCESS_STOCK,))
        self.assertEqual(unknown_processes, expected_all)

    def test_updates_load_excel_runs_in_background_with_overlay_and_double_click_guard(self) -> None:
        from futonhub.ui.erp.prototype import FutonHubErpPrototype

        source = inspect.getsource(FutonHubErpPrototype._build_updates)
        self.assertIn('if state.get("applying") or state.get("loading"):', source)
        self.assertIn('self._show_working_overlay("Cargando Excel"', source)
        self.assertIn("set_loading(True, loading_message)", source)
        self.assertIn("set_loading(False)", source)
        self.assertIn("Cargando y validando stock...", source)
        self.assertIn("Preparando vista previa...", source)
        self.assertIn("preview_updates_from_excel(session, process, path)", source)


if __name__ == "__main__":
    unittest.main()
